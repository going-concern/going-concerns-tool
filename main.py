from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src.control_model import ControlEvaluation, ControlTestRequest, evaluate_control


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
TITLE_FILL = PatternFill("solid", fgColor="E2EFDA")
BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


def _apply_header(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=True)
    cell.fill = HEADER_FILL
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center")
    return cell


def _apply_title(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=True, size=14)
    cell.fill = TITLE_FILL
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="left")
    return cell


def _add_summary_sheet(wb: Workbook, control: ControlTestRequest, evaluation: ControlEvaluation) -> None:
    ws = wb.active
    ws.title = "TOE Summary"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Control Testing Working Paper"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].fill = TITLE_FILL

    rows = [
        ["Control ID", control.control_id],
        ["Control Name", control.control_name],
        ["Objective", control.objective],
        ["Testing Period", control.testing_period],
        ["Population Definition", control.population_definition],
        ["Sample Size", evaluation.total_samples],
        ["Pass Count", evaluation.pass_count],
        ["Exception Count", evaluation.fail_count],
        ["Inconclusive Count", evaluation.inconclusive_count],
        ["Observed Deviation Rate", f"{evaluation.observed_deviation_rate}%" if evaluation.observed_deviation_rate is not None else "Not calculated"],
        ["Tolerable Deviation Rate", f"{evaluation.tolerable_deviation_rate}%" if evaluation.tolerable_deviation_rate is not None else "Not configured"],
        ["Calculation Summary", evaluation.summary],
        ["Auditor Evaluation Required", "Yes" if evaluation.auditor_evaluation_required else "No"],
    ]

    for idx, row in enumerate(rows, start=3):
        ws.cell(row=idx, column=1, value=row[0])
        ws.cell(row=idx, column=2, value=row[1])
        ws.cell(row=idx, column=1).font = Font(bold=True)
        ws.cell(row=idx, column=1).fill = HEADER_FILL
        ws.cell(row=idx, column=1).border = BORDER
        ws.cell(row=idx, column=2).border = BORDER
        ws.cell(row=idx, column=2).alignment = Alignment(wrap_text=True)

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90


def _add_design_sheet(wb: Workbook, control: ControlTestRequest) -> None:
    ws = wb.create_sheet("Design & Implementation")
    ws.freeze_panes = "A2"
    _apply_title(ws, 1, 1, "Design and Implementation Assessment")

    rows = [
        ["Control Description", control.control_description],
        ["Relevant Risk", control.relevant_risk],
        ["Relevant Assertions", control.relevant_assertions],
        ["Financial Statement Area", control.financial_statement_area],
        ["Process", control.process],
        ["Control Owner", control.control_owner],
        ["D&I Conclusion", control.di_conclusion],
        ["D&I Rationale", control.di_rationale],
    ]

    for idx, row in enumerate(rows, start=3):
        _apply_header(ws, idx, 1, row[0])
        ws.cell(row=idx, column=2, value=row[1])
        ws.cell(row=idx, column=2).border = BORDER
        ws.cell(row=idx, column=2).alignment = Alignment(wrap_text=True)

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100


def _add_sample_results_sheet(wb: Workbook, control: ControlTestRequest) -> None:
    ws = wb.create_sheet("Sample Results")
    ws.freeze_panes = "A2"
    _apply_title(ws, 1, 1, "Sample Testing")

    headers = ["Sample ID", "Population Record ID", "Sample Number", "Result", "Exception Indicator", "Exception Type", "Evidence IDs", "Tester", "Testing Date", "Notes"]
    for col, value in enumerate(headers, start=1):
        _apply_header(ws, 3, col, value)

    for idx, sample in enumerate(control.sample_results, start=4):
        ws.cell(row=idx, column=1, value=sample.sample_id)
        ws.cell(row=idx, column=2, value=sample.population_record_id)
        ws.cell(row=idx, column=3, value=sample.sample_number)
        ws.cell(row=idx, column=4, value=sample.test_result)
        ws.cell(row=idx, column=5, value="Yes" if sample.exception_indicator else "No")
        ws.cell(row=idx, column=6, value=sample.exception_type)
        ws.cell(row=idx, column=7, value=", ".join(sample.evidence_ids))
        ws.cell(row=idx, column=8, value=sample.tester)
        ws.cell(row=idx, column=9, value=sample.testing_date)
        ws.cell(row=idx, column=10, value=sample.notes)

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 24
    ws.column_dimensions["G"].width = 28
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["I"].width = 18
    ws.column_dimensions["J"].width = 36


def _add_review_sheet(wb: Workbook, control: ControlTestRequest, evaluation: ControlEvaluation) -> None:
    ws = wb.create_sheet("Review & Notes")
    _apply_title(ws, 1, 1, "Reviewer Notes and Final Conclusion")

    ws["A3"] = "Calculation Summary"
    ws["A3"].font = Font(bold=True)
    ws["B3"] = evaluation.summary
    ws["B3"].alignment = Alignment(wrap_text=True)

    ws["A5"] = "Auditor Evaluation Required"
    ws["A5"].font = Font(bold=True)
    ws["B5"] = "Yes" if evaluation.auditor_evaluation_required else "No"

    ws["A7"] = "Reviewer Notes"
    ws["A7"].font = Font(bold=True)
    ws["B7"] = control.reviewer_notes or "No reviewer notes recorded."
    ws["B7"].alignment = Alignment(wrap_text=True)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90


def generate_toe_workbook(control: ControlTestRequest, output_path: str) -> None:
    evaluation = evaluate_control(control)
    wb = Workbook()
    _add_summary_sheet(wb, control, evaluation)
    _add_design_sheet(wb, control)
    _add_sample_results_sheet(wb, control)
    _add_review_sheet(wb, control, evaluation)
    wb.save(output_path)


__all__ = ["generate_toe_workbook"]
